#!/usr/bin/env python3
"""Parse spreadsheets and delimited text files into Markdown and structured rows."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


CSV_EXTENSIONS = {".csv", ".tsv"}
EXCEL_EXTENSIONS = {".xls", ".xlsx", ".xlsm", ".xlsb", ".ods"}
MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PACKAGE_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _row_limit(max_rows: int | None) -> int | None:
    if max_rows is None or max_rows <= 0:
        return None
    return max_rows + 1


def _cell(value: object, max_len: int = 120) -> str:
    if value is None:
        text = ""
    else:
        text = str(value)
    text = text.replace("\r", " ").replace("\n", " ").strip()
    if len(text) > max_len:
        text = text[: max_len - 3] + "..."
    return text.replace("|", "\\|")


def _rows_to_markdown(rows: list[list[object]], title: str) -> str:
    if not rows:
        return f"## {title}\n\n(empty)\n"
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = [_cell(value) or f"Column {index + 1}" for index, value in enumerate(normalized[0])]
    body = normalized[1:]
    lines = [f"## {title}", "", "| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(_cell(value) for value in row) + " |" for row in body)
    lines.append("")
    return "\n".join(lines)


def _column_index(cell_ref: str) -> int:
    letters = re.sub(r"[^A-Z]", "", cell_ref.upper())
    index = 0
    for letter in letters:
        index = index * 26 + (ord(letter) - ord("A") + 1)
    return max(index, 1)


def _load_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    try:
        payload = archive.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    root = ET.fromstring(payload)
    values = []
    for item in root.findall(MAIN_NS + "si"):
        parts = []
        for text_node in item.iter(MAIN_NS + "t"):
            parts.append(text_node.text or "")
        values.append("".join(parts))
    return values


def _sheet_paths(archive: zipfile.ZipFile) -> list[tuple[str, str]]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {}
    for rel in rels.findall(PACKAGE_REL_NS + "Relationship"):
        target = rel.attrib.get("Target", "")
        if target.startswith("/"):
            path = target.lstrip("/")
        else:
            path = "xl/" + target.lstrip("/")
        rel_targets[rel.attrib.get("Id", "")] = path

    sheets = []
    sheets_node = workbook.find(MAIN_NS + "sheets")
    if sheets_node is None:
        return sheets
    for sheet_node in sheets_node.findall(MAIN_NS + "sheet"):
        name = sheet_node.attrib.get("name", "Sheet")
        rid = sheet_node.attrib.get(REL_NS + "id", "")
        target = rel_targets.get(rid)
        if target:
            sheets.append((name, target))
    return sheets


def _cell_text(cell: ET.Element, shared_strings: list[str]) -> str:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        inline = cell.find(MAIN_NS + "is")
        if inline is None:
            return ""
        return "".join(text_node.text or "" for text_node in inline.iter(MAIN_NS + "t"))

    value = cell.find(MAIN_NS + "v")
    if value is None or value.text is None:
        formula = cell.find(MAIN_NS + "f")
        return f"={formula.text}" if formula is not None and formula.text else ""

    raw = value.text
    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except (ValueError, IndexError):
            return raw
    if cell_type == "b":
        return "TRUE" if raw == "1" else "FALSE"
    return raw


def _parse_sheet_xml(
    archive: zipfile.ZipFile,
    sheet_path: str,
    shared_strings: list[str],
    max_rows: int | None,
) -> tuple[list[list[object]], bool, int | None, int | None]:
    row_limit = _row_limit(max_rows)
    rows: list[list[object]] = []
    truncated = False
    max_row_seen: int | None = None
    max_col_seen: int | None = None

    with archive.open(sheet_path) as handle:
        for event, elem in ET.iterparse(handle, events=("end",)):
            if elem.tag != MAIN_NS + "row":
                continue
            if row_limit is not None and len(rows) >= row_limit:
                truncated = True
                elem.clear()
                break

            row_index_text = elem.attrib.get("r")
            if row_index_text and row_index_text.isdigit():
                max_row_seen = max(max_row_seen or 0, int(row_index_text))

            row_values: list[object] = []
            for cell in elem.findall(MAIN_NS + "c"):
                ref = cell.attrib.get("r", "")
                col_index = _column_index(ref)
                max_col_seen = max(max_col_seen or 0, col_index)
                while len(row_values) < col_index - 1:
                    row_values.append("")
                row_values.append(_cell_text(cell, shared_strings))
            rows.append(row_values)
            elem.clear()

    return rows, truncated, max_row_seen, max_col_seen


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def _parse_csv(path: Path, max_rows: int | None) -> dict:
    text = _read_text(path)
    sample = text[:4096]
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=[",", "\t", ";", "|"])
        delimiter = dialect.delimiter
    except csv.Error:
        pass

    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    row_limit = _row_limit(max_rows)
    rows = []
    truncated = False
    for index, row in enumerate(reader):
        if row_limit is not None and index >= row_limit:
            truncated = True
            break
        rows.append(row)

    warnings = []
    if truncated:
        warnings.append(f"CSV preview limited to {max_rows} data rows.")

    return {
        "parser": "csv",
        "file_type": path.suffix.lower().lstrip(".") or "csv",
        "source": str(path),
        "sheets": [
            {
                "name": path.name,
                "rows": rows,
                "row_count_preview": len(rows),
                "truncated": truncated,
                "max_rows_requested": max_rows,
            }
        ],
        "warnings": warnings,
        "markdown": f"# {path.name}\n\n" + _rows_to_markdown(rows, path.name),
    }


def _parse_with_openpyxl(path: Path, max_rows: int | None, sheet: str | None, data_only: bool) -> dict:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
    sheet_names = [sheet] if sheet else workbook.sheetnames
    payloads = []
    all_warnings = []
    markdown_parts = [f"# {path.name}", "", "Parser: `openpyxl`", ""]

    for sheet_name in sheet_names:
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(f"Sheet not found: {sheet_name}. Available: {', '.join(workbook.sheetnames)}")
        ws = workbook[sheet_name]
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        row_limit = _row_limit(max_rows)
        rows = []
        truncated = False
        for index, row in enumerate(ws.iter_rows(values_only=True)):
            if row_limit is not None and index >= row_limit:
                truncated = True
                break
            rows.append(list(row))
        warnings = []
        if truncated:
            warnings.append(f"Sheet {sheet_name} preview limited to {max_rows} data rows.")
        if len(rows) <= 1 and (ws.max_row or 0) > len(rows):
            limit_hint = (
                "Increase --max-rows or inspect workbook structure."
                if max_rows is not None and max_rows > 0
                else "Inspect workbook structure."
            )
            warnings.append(
                f"Sheet {sheet_name} produced only {len(rows)} preview row(s), but workbook reports max_row={ws.max_row}. "
                f"{limit_hint}"
            )
        all_warnings.extend(warnings)
        payloads.append(
            {
                "name": sheet_name,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "rows": rows,
                "row_count_preview": len(rows),
                "truncated": truncated,
                "max_rows_requested": max_rows,
                "warnings": warnings,
            }
        )
        markdown_parts.append(_rows_to_markdown(rows, sheet_name))

    return {
        "parser": "openpyxl",
        "file_type": path.suffix.lower().lstrip("."),
        "source": str(path),
        "sheets": payloads,
        "warnings": all_warnings,
        "markdown": "\n".join(markdown_parts).strip(),
    }


def _parse_xlsx_with_zip_xml(path: Path, max_rows: int | None, sheet: str | None) -> dict:
    payloads = []
    all_warnings = []
    markdown_parts = [f"# {path.name}", "", "Parser: `zip-xml-xlsx`", ""]

    with zipfile.ZipFile(path) as archive:
        shared_strings = _load_shared_strings(archive)
        sheets = _sheet_paths(archive)
        if sheet:
            sheets = [(name, target) for name, target in sheets if name == sheet]
            if not sheets:
                available = ", ".join(name for name, _ in _sheet_paths(archive))
                raise RuntimeError(f"Sheet not found: {sheet}. Available: {available}")

        for sheet_name, sheet_path in sheets:
            rows, truncated, max_row_seen, max_col_seen = _parse_sheet_xml(
                archive,
                sheet_path,
                shared_strings,
                max_rows,
            )
            warnings = []
            if truncated:
                warnings.append(f"Sheet {sheet_name} preview limited to {max_rows} data rows.")
            if len(rows) <= 1 and (max_row_seen or 0) > len(rows):
                warnings.append(
                    f"Sheet {sheet_name} produced only {len(rows)} row(s), but XML reports row {max_row_seen}. "
                    "Inspect workbook structure."
                )
            all_warnings.extend(warnings)
            payloads.append(
                {
                    "name": sheet_name,
                    "max_row": None if truncated else max_row_seen,
                    "last_row_read": max_row_seen,
                    "max_column": max_col_seen,
                    "rows": rows,
                    "row_count_preview": len(rows),
                    "truncated": truncated,
                    "max_rows_requested": max_rows,
                    "warnings": warnings,
                }
            )
            markdown_parts.append(_rows_to_markdown(rows, sheet_name))

    return {
        "parser": "zip-xml-xlsx",
        "file_type": path.suffix.lower().lstrip("."),
        "source": str(path),
        "sheets": payloads,
        "warnings": all_warnings,
        "markdown": "\n".join(markdown_parts).strip(),
    }


def _preview_row_count(result: dict) -> int:
    sheets = result.get("sheets", [])
    if not sheets:
        return 0
    return max(int(sheet.get("row_count_preview") or 0) for sheet in sheets)


def _dataframe_rows(frame) -> list[list[object]]:
    columns = list(frame.columns)
    values = frame.where(frame.notna(), "").values.tolist()
    return [columns] + values


def _parse_with_pandas(path: Path, max_rows: int | None, sheet: str | None) -> dict:
    import pandas as pd

    sheet_name: str | int | None = sheet if sheet else None
    nrows = None if max_rows is None or max_rows <= 0 else max_rows
    loaded = pd.read_excel(path, sheet_name=sheet_name, nrows=nrows)
    if not isinstance(loaded, dict):
        loaded = {sheet or "Sheet1": loaded}

    payloads = []
    markdown_parts = [f"# {path.name}", "", "Parser: `pandas.read_excel`", ""]
    for sheet_name, frame in loaded.items():
        rows = _dataframe_rows(frame)
        payloads.append(
            {
                "name": str(sheet_name),
                "columns": [str(col) for col in frame.columns],
                "rows": rows,
                "row_count_preview": len(rows),
                "truncated": nrows is not None and len(frame.index) >= nrows,
                "max_rows_requested": max_rows,
            }
        )
        markdown_parts.append(_rows_to_markdown(rows, str(sheet_name)))

    return {
        "parser": "pandas",
        "file_type": path.suffix.lower().lstrip("."),
        "source": str(path),
        "sheets": payloads,
        "warnings": [],
        "markdown": "\n".join(markdown_parts).strip(),
    }


def _parse_with_xlrd(path: Path, max_rows: int | None, sheet: str | None) -> dict:
    import xlrd

    workbook = xlrd.open_workbook(str(path), on_demand=True)
    sheet_names = [sheet] if sheet else workbook.sheet_names()
    payloads = []
    markdown_parts = [f"# {path.name}", "", "Parser: `xlrd`", ""]

    for sheet_name in sheet_names:
        if sheet_name not in workbook.sheet_names():
            raise RuntimeError(f"Sheet not found: {sheet_name}. Available: {', '.join(workbook.sheet_names())}")
        ws = workbook.sheet_by_name(sheet_name)
        preview_rows = ws.nrows if max_rows is None or max_rows <= 0 else min(ws.nrows, max_rows + 1)
        rows = [ws.row_values(index) for index in range(preview_rows)]
        payloads.append(
            {
                "name": sheet_name,
                "max_row": ws.nrows,
                "max_column": ws.ncols,
                "rows": rows,
                "row_count_preview": len(rows),
                "truncated": preview_rows < ws.nrows,
                "max_rows_requested": max_rows,
            }
        )
        markdown_parts.append(_rows_to_markdown(rows, sheet_name))

    return {
        "parser": "xlrd",
        "file_type": path.suffix.lower().lstrip("."),
        "source": str(path),
        "sheets": payloads,
        "warnings": [],
        "markdown": "\n".join(markdown_parts).strip(),
    }


def parse_spreadsheet(
    path: Path,
    max_rows: int = 0,
    sheet: str | None = None,
    data_only: bool = False,
) -> dict:
    suffix = path.suffix.lower()
    if suffix in CSV_EXTENSIONS:
        return _parse_csv(path, max_rows=max_rows)
    if suffix not in EXCEL_EXTENSIONS:
        raise RuntimeError(f"Unsupported spreadsheet type: {suffix}")

    errors: list[str] = []
    if suffix in {".xlsx", ".xlsm"}:
        try:
            result = _parse_with_openpyxl(path, max_rows=max_rows, sheet=sheet, data_only=data_only)
            if _preview_row_count(result) <= 1:
                try:
                    fallback = _parse_xlsx_with_zip_xml(path, max_rows=max_rows, sheet=sheet)
                    if _preview_row_count(fallback) > _preview_row_count(result):
                        return fallback
                except Exception as exc:
                    errors.append(f"zip-xml-xlsx verification failed: {exc}")
            return result
        except ModuleNotFoundError as exc:
            errors.append(f"openpyxl missing ({exc.name}). Install with: python3 -m pip install openpyxl")
        except Exception as exc:
            errors.append(f"openpyxl failed: {exc}")

        try:
            return _parse_xlsx_with_zip_xml(path, max_rows=max_rows, sheet=sheet)
        except Exception as exc:
            errors.append(f"zip-xml-xlsx failed: {exc}")

    try:
        return _parse_with_pandas(path, max_rows=max_rows, sheet=sheet)
    except ModuleNotFoundError as exc:
        errors.append(
            f"pandas or engine missing ({exc.name}). Install with: "
            "python3 -m pip install pandas openpyxl xlrd pyxlsb python-calamine odfpy"
        )
    except Exception as exc:
        errors.append(f"pandas failed: {exc}")

    if suffix == ".xls":
        try:
            return _parse_with_xlrd(path, max_rows=max_rows, sheet=sheet)
        except ModuleNotFoundError as exc:
            errors.append(f"xlrd missing ({exc.name}). Install with: python3 -m pip install xlrd")
        except Exception as exc:
            errors.append(f"xlrd failed: {exc}")

    raise RuntimeError("No spreadsheet parser succeeded. " + " | ".join(errors))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Parse spreadsheet files into Markdown/JSON.")
    parser.add_argument("input")
    parser.add_argument("-o", "--output", help="Markdown output path")
    parser.add_argument("--json", help="JSON output path")
    parser.add_argument("--sheet", help="Sheet name")
    parser.add_argument(
        "--max-rows",
        type=int,
        default=0,
        help="Maximum data rows per sheet. Default 0 means no limit.",
    )
    parser.add_argument("--all-rows", action="store_true", help="Read all rows. Equivalent to --max-rows 0.")
    parser.add_argument("--data-only", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    path = Path(args.input).expanduser().resolve()
    max_rows = 0 if args.all_rows else args.max_rows
    try:
        result = parse_spreadsheet(path, max_rows=max_rows, sheet=args.sheet, data_only=args.data_only)
    except Exception as exc:
        print(f"Spreadsheet parse failed: {exc}", file=sys.stderr)
        return 1

    if args.output:
        Path(args.output).expanduser().resolve().write_text(result["markdown"] + "\n", encoding="utf-8")
    else:
        print(result["markdown"])
    if args.json:
        Path(args.json).expanduser().resolve().write_text(
            json.dumps({k: v for k, v in result.items() if k != "markdown"}, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

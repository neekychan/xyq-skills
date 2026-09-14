#!/usr/bin/env python3
"""Parse DOCX and PPTX files with optional package-backed or ZIP XML fallbacks."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _truncate(text: str, max_chars: int | None) -> str:
    if max_chars is None or max_chars <= 0 or len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[truncated]\n"


def _docx_with_python_docx(path: Path, max_chars: int | None) -> dict:
    import docx

    document = docx.Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)

    markdown_parts = [f"# {path.name}", "", "Parser: `python-docx`", ""]
    markdown_parts.extend(paragraphs)
    for index, rows in enumerate(tables, start=1):
        markdown_parts.extend(["", f"## Table {index}", "", _rows_to_markdown(rows), ""])
    return {
        "parser": "python-docx",
        "file_type": "docx",
        "source": str(path),
        "paragraphs": paragraphs,
        "tables": tables,
        "warnings": [],
        "markdown": _truncate("\n\n".join(markdown_parts).strip(), max_chars),
    }


def _text_from_xml_element(element: ET.Element, tag: str) -> str:
    return "".join(node.text or "" for node in element.iter(tag)).strip()


def _docx_with_zip_xml(path: Path, max_chars: int | None) -> dict:
    paragraphs = []
    tables = []
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    body = root.find(f"{W_NS}body")
    if body is not None:
        for child in body:
            if child.tag == f"{W_NS}p":
                text = _text_from_xml_element(child, f"{W_NS}t")
                if text:
                    paragraphs.append(text)
            elif child.tag == f"{W_NS}tbl":
                rows = []
                for tr in child.iter(f"{W_NS}tr"):
                    row = []
                    for tc in tr.iter(f"{W_NS}tc"):
                        row.append(_text_from_xml_element(tc, f"{W_NS}t"))
                    if row:
                        rows.append(row)
                if rows:
                    tables.append(rows)

    markdown_parts = [f"# {path.name}", "", "Parser: `zip-xml-docx`", ""]
    markdown_parts.extend(paragraphs)
    for index, rows in enumerate(tables, start=1):
        markdown_parts.extend(["", f"## Table {index}", "", _rows_to_markdown(rows), ""])
    return {
        "parser": "zip-xml-docx",
        "file_type": "docx",
        "source": str(path),
        "paragraphs": paragraphs,
        "tables": tables,
        "warnings": ["Used basic ZIP XML fallback; comments, footnotes, and advanced formatting may be omitted."],
        "markdown": _truncate("\n\n".join(markdown_parts).strip(), max_chars),
    }


def _rows_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = [cell.replace("|", "\\|") for cell in normalized[0]]
    body = normalized[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |" for row in body)
    return "\n".join(lines)


def parse_docx(path: Path, max_chars: int | None = 0) -> dict:
    try:
        return _docx_with_python_docx(path, max_chars=max_chars)
    except ModuleNotFoundError:
        return _docx_with_zip_xml(path, max_chars=max_chars)


def _pptx_with_python_pptx(path: Path, max_chars: int | None) -> dict:
    import pptx

    deck = pptx.Presentation(str(path))
    slides = []
    markdown_parts = [f"# {path.name}", "", "Parser: `python-pptx`", ""]
    for slide_index, slide in enumerate(deck.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        slides.append({"slide": slide_index, "texts": texts})
        markdown_parts.extend([f"## Slide {slide_index}", ""])
        markdown_parts.extend(texts or ["(no text)"])
        markdown_parts.append("")
    return {
        "parser": "python-pptx",
        "file_type": "pptx",
        "source": str(path),
        "slides": slides,
        "warnings": [],
        "markdown": _truncate("\n".join(markdown_parts).strip(), max_chars),
    }


def _pptx_with_zip_xml(path: Path, max_chars: int | None) -> dict:
    slides = []
    markdown_parts = [f"# {path.name}", "", "Parser: `zip-xml-pptx`", ""]
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            name for name in archive.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)
        )
        for slide_index, name in enumerate(slide_names, start=1):
            root = ET.fromstring(archive.read(name))
            texts = [node.text.strip() for node in root.iter(f"{A_NS}t") if node.text and node.text.strip()]
            slides.append({"slide": slide_index, "texts": texts})
            markdown_parts.extend([f"## Slide {slide_index}", ""])
            markdown_parts.extend(texts or ["(no text)"])
            markdown_parts.append("")
    return {
        "parser": "zip-xml-pptx",
        "file_type": "pptx",
        "source": str(path),
        "slides": slides,
        "warnings": ["Used basic ZIP XML fallback; speaker notes and advanced formatting may be omitted."],
        "markdown": _truncate("\n".join(markdown_parts).strip(), max_chars),
    }


def parse_pptx(path: Path, max_chars: int | None = 0) -> dict:
    try:
        return _pptx_with_python_pptx(path, max_chars=max_chars)
    except ModuleNotFoundError:
        return _pptx_with_zip_xml(path, max_chars=max_chars)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Parse DOCX or PPTX into Markdown/JSON.")
    parser.add_argument("input")
    parser.add_argument("-o", "--output", help="Markdown output path")
    parser.add_argument("--json", help="JSON output path")
    parser.add_argument("--max-chars", type=int, default=0, help="Maximum Markdown characters. Default 0 means no limit.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    path = Path(args.input).expanduser().resolve()
    try:
        if path.suffix.lower() == ".docx":
            result = parse_docx(path, max_chars=args.max_chars)
        elif path.suffix.lower() == ".pptx":
            result = parse_pptx(path, max_chars=args.max_chars)
        else:
            raise RuntimeError("Only .docx and .pptx are supported directly. Convert .doc/.ppt first.")
    except Exception as exc:
        print(f"Office parse failed: {exc}", file=sys.stderr)
        return 1

    if args.output:
        Path(args.output).expanduser().resolve().write_text(result["markdown"] + "\n", encoding="utf-8")
    else:
        print(result["markdown"])
    if args.json:
        Path(args.json).expanduser().resolve().write_text(
            json.dumps({k: v for k, v in result.items() if k != "markdown"}, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
